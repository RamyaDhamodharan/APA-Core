"""Idempotently seed vendors, purchase orders, and goods receipts from CSV/XLSX."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


def _rows(path: Path, file_format: str) -> list[dict[str, Any]]:
    selected = file_format.casefold()
    if selected == "auto":
        selected = path.suffix.lstrip(".").casefold()
    if selected == "csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return [
                {str(key).strip(): value for key, value in row.items() if key}
                for row in csv.DictReader(stream)
            ]
    if selected in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [
            str(value).strip() if value is not None else ""
            for value in next(iterator, ())
        ]
        rows = [
            {
                header: value
                for header, value in zip(headers, values, strict=False)
                if header
            }
            for values in iterator
        ]
        workbook.close()
        return rows
    raise ValueError(f"Unsupported file format for {path.name}: {file_format}")


def _value(row: dict[str, Any], *names: str, default=None):
    normalized = {str(key).strip().casefold(): value for key, value in row.items()}
    for name in names:
        value = normalized.get(name.casefold())
        if value not in (None, ""):
            return value
    return default


def _decimal(value, *, default=None):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid decimal value: {value}") from exc


def _date(value, *, default=None):
    if value in (None, ""):
        return default
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Invalid date value: {value}")


def _counted(db, items, create_one):
    summary = {"created": 0, "skipped": 0, "errors": []}
    for label, item in items:
        try:
            outcome = create_one(item)
            summary[outcome] += 1
        except (
            Exception
        ) as exc:  # Keep one malformed source row from stopping the import.
            db.rollback()
            summary["errors"].append({"item": label, "message": str(exc)[:300]})
    return summary


def import_vendors(db, rows):
    from sqlalchemy import func

    from src.models.vendor_models import Vendor

    def create(row):
        name = str(_value(row, "name", "vendor", "vendor_name", default="")).strip()
        if not name:
            raise ValueError("Vendor name is required")
        existing = (
            db.query(Vendor).filter(func.lower(Vendor.name) == name.casefold()).first()
        )
        if existing:
            return "skipped"
        db.add(
            Vendor(
                name=name,
                tax_id=_value(row, "tax_id", "tax_id_number"),
                default_tax_rate=_decimal(_value(row, "default_tax_rate", "tax_rate")),
                contact_name=_value(row, "contact_name"),
                contact_email=_value(row, "contact_email", "email"),
                contact_phone=_value(row, "contact_phone", "phone"),
                is_active=True,
            )
        )
        db.commit()
        return "created"

    return _counted(
        db,
        (
            (str(_value(row, "name", "vendor", default=index + 1)), row)
            for index, row in enumerate(rows)
        ),
        create,
    )


def import_purchase_orders(db, rows):
    from src.models.purchase_order_models import PurchaseOrder, PurchaseOrderLine
    from src.models.vendor_models import Vendor

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        number = str(
            _value(row, "po_number", "purchase_order", "order_number", default="")
        ).strip()
        if not number:
            grouped[f"__invalid_{len(grouped)}"].append(row)
        else:
            grouped[number].append(row)

    def create(group):
        number, lines = group
        if number.startswith("__invalid_"):
            raise ValueError("PO number is required")
        if db.query(PurchaseOrder).filter_by(po_number=number, is_current=True).first():
            return "skipped"
        first = lines[0]
        vendor_name = str(_value(first, "vendor", "vendor_name", default="")).strip()
        vendor = (
            db.query(Vendor).filter(Vendor.name.ilike(vendor_name)).first()
            if vendor_name
            else None
        )
        if vendor is None:
            raise ValueError(f"Vendor not found: {vendor_name or '(missing)'}")
        po_lines = []
        for line in lines:
            qty = _decimal(_value(line, "qty", "quantity"))
            price = _decimal(_value(line, "unit_price", "price"))
            amount = _decimal(
                _value(line, "line_amount", "amount"),
                default=qty * price if qty is not None and price is not None else None,
            )
            if not any((qty, price, amount, _value(line, "item_code", "description"))):
                continue
            po_lines.append(
                PurchaseOrderLine(
                    item_code=_value(line, "item_code", "sku"),
                    description=_value(line, "description", "item_description"),
                    uom=_value(line, "uom", "unit"),
                    qty=qty,
                    unit_price=price,
                    line_amount=amount,
                )
            )
        subtotal = _decimal(
            _value(first, "subtotal"),
            default=sum(
                (line.line_amount or Decimal(0) for line in po_lines), Decimal(0)
            ),
        )
        tax = _decimal(_value(first, "tax_amount", "tax"), default=Decimal(0))
        order = PurchaseOrder(
            po_number=number,
            vendor_id=vendor.id,
            po_date=_date(_value(first, "po_date", "date")),
            currency=str(_value(first, "currency", default="INR")).upper(),
            subtotal=subtotal,
            tax_amount=tax,
            tax_rate=_decimal(_value(first, "tax_rate")),
            total_amount=_decimal(
                _value(first, "total_amount", "total"), default=subtotal + tax
            ),
            is_current=True,
            lines=po_lines,
        )
        db.add(order)
        db.commit()
        return "created"

    entries = [(number, (number, lines)) for number, lines in grouped.items()]
    return _counted(db, entries, create)


def import_goods_receipts(db, rows):
    from src.models.goods_receipt_models import GoodsReceipt, GoodsReceiptLine
    from src.models.purchase_order_models import PurchaseOrder

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        number = str(
            _value(row, "gr_number", "receipt_number", "goods_receipt", default="")
        ).strip()
        grouped[number or f"__invalid_{len(grouped)}"].append(row)

    def create(group):
        number, lines = group
        if number.startswith("__invalid_"):
            raise ValueError("Goods receipt number is required")
        if db.query(GoodsReceipt).filter_by(gr_number=number).first():
            return "skipped"
        first = lines[0]
        po_number = str(
            _value(first, "po_number", "purchase_order", "order_number", default="")
        ).strip()
        po = (
            db.query(PurchaseOrder)
            .filter_by(po_number=po_number, is_current=True)
            .first()
        )
        if po is None:
            raise ValueError(f"Purchase order not found: {po_number or '(missing)'}")
        po_lines = {str(line.item_code or "").casefold(): line for line in po.lines}
        receipt_lines = []
        for row in lines:
            item_code = str(_value(row, "item_code", "sku", default="")).strip()
            if not item_code:
                continue
            po_line = po_lines.get(item_code.casefold())
            if po_line is None:
                raise ValueError(f"PO line not found for item {item_code}")
            received = _decimal(_value(row, "received_qty", "quantity", "qty"))
            accepted = _decimal(_value(row, "accepted_qty"), default=received)
            rejected = _decimal(_value(row, "rejected_qty"), default=Decimal(0))
            if received is None or accepted + rejected != received:
                raise ValueError(
                    f"Invalid received/accepted/rejected quantity for {item_code}"
                )
            receipt_lines.append(
                GoodsReceiptLine(
                    po_line_id=po_line.id,
                    item_code=po_line.item_code,
                    description=po_line.description,
                    uom=po_line.uom,
                    received_qty=received,
                    accepted_qty=accepted,
                    rejected_qty=rejected,
                )
            )
        if not receipt_lines:
            raise ValueError("Goods receipt requires at least one item line")
        db.add(
            GoodsReceipt(
                gr_number=number,
                po_id=po.id,
                vendor_id=po.vendor_id,
                receipt_date=_date(
                    _value(first, "receipt_date", "date"), default=date.today()
                ),
                lines=receipt_lines,
            )
        )
        db.commit()
        return "created"

    entries = [(number, (number, lines)) for number, lines in grouped.items()]
    return _counted(db, entries, create)


def _print_summary(label, summary):
    print(
        f"{label}: created={summary['created']} skipped={summary['skipped']} "
        f"errors={len(summary['errors'])}"
    )
    for error in summary["errors"]:
        print(f"  ERROR {error['item']}: {error['message']}")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--folder", type=Path, required=True)
    parser.add_argument("--vendors", default="vendors.csv")
    parser.add_argument("--purchase-orders", default="purchase_orders.csv")
    parser.add_argument("--goods-receipts", default="goods_receipts.csv")
    for key in ("vendors", "purchase_orders", "goods_receipts"):
        parser.add_argument(
            f"--{key.replace('_', '-')}-format",
            choices=("auto", "csv", "xlsx", "xlsm"),
            default="auto",
        )
    args = parser.parse_args(argv)

    from src.db.session_db import SessionLocal

    db = SessionLocal()
    try:
        summary = {}
        for label, filename, file_format, importer in (
            ("Vendors", args.vendors, args.vendors_format, import_vendors),
            (
                "Purchase orders",
                args.purchase_orders,
                args.purchase_orders_format,
                import_purchase_orders,
            ),
            (
                "Goods receipts",
                args.goods_receipts,
                args.goods_receipts_format,
                import_goods_receipts,
            ),
        ):
            path = args.folder / filename
            try:
                summary[label] = importer(db, _rows(path, file_format))
            except Exception as exc:
                db.rollback()
                summary[label] = {
                    "created": 0,
                    "skipped": 0,
                    "errors": [{"item": str(path), "message": str(exc)[:300]}],
                }
        for label, result in summary.items():
            _print_summary(label, result)
        return int(any(result["errors"] for result in summary.values()))
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
